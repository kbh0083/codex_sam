from __future__ import annotations

import argparse
import base64
import csv
import hashlib
import json
import os
import re
import sys
import time
import traceback
import unicodedata
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import fitz
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from PIL import Image, ImageDraw, ImageFont
from weasyprint import HTML


ROOT = Path(__file__).resolve().parents[2]
DOCUMENT_DIR = ROOT / "document"
ANSWER_DIR = ROOT / "거래처별_문서_정답"
REPORT_ROOT = ROOT / "tests" / "variable_data" / "report"
PROMPT_DIR = ROOT / "tests" / "variable_data" / "prompts"
KST = ZoneInfo("Asia/Seoul")

DEFAULT_BASE_URL = "https://api.novita.ai/openai"
DEFAULT_MODEL = "qwen/qwen3.5-397b-a17b"

CSV_FIELDS = [
    "fund_code",
    "fund_name",
    "base_date",
    "t_day",
    "transfer_amount",
    "settle_class",
    "order_type",
]

SETTLE_TO_CODE = {
    "1": "1",
    "2": "2",
    "청구": "1",
    "확정": "2",
    "예정": "1",
    "펀드매입": "1",
    "펀드환매": "2",
    "PENDING": "1",
    "CONFIRMED": "2",
}
ORDER_TO_CODE = {
    "1": "1",
    "3": "3",
    "4": "1",
    "해지": "1",
    "설정": "3",
    "매도": "1",
    "매수": "3",
    "RED": "1",
    "SUB": "3",
}


class RunBlockedError(RuntimeError):
    """External API or runtime condition blocked a test case execution."""


@dataclass(frozen=True)
class Profile:
    name: str
    enable_thinking: bool
    temperature: float
    top_p: float
    top_k: int
    min_p: float
    presence_penalty: float
    repetition_penalty: float
    max_tokens: int


@dataclass(frozen=True)
class CaseSpec:
    case_id: str
    company: str
    source_name: str
    answer_dir_name: str
    answer_json_name: str
    answer_csv_name: str
    expected_orders: int
    base_date: str
    kind: str
    strategy: str
    compact_fund_name: bool = False
    sheet_name: str | None = None
    xlsx_range: str | None = None


@dataclass(frozen=True)
class RenderCell:
    row: int
    col: int
    rowspan: int
    colspan: int
    text: str


PROFILES: dict[str, Profile] = {
    "non_thinking": Profile(
        name="non_thinking",
        enable_thinking=False,
        temperature=0,
        top_p=1.0,
        top_k=20,
        min_p=0.0,
        presence_penalty=0.0,
        repetition_penalty=1.0,
        max_tokens=16384,
    ),
    "thinking": Profile(
        name="thinking",
        enable_thinking=True,
        temperature=0.6,
        top_p=0.95,
        top_k=20,
        min_p=0.0,
        presence_penalty=0.0,
        repetition_penalty=1.0,
        max_tokens=32768,
    ),
}

CASE_SPECS: list[CaseSpec] = [
    CaseSpec(
        case_id="33_운용지시서_KDB생명_20260427_eml",
        company="KDB생명",
        source_name="운용지시서(KDB생명)_20260427.eml",
        answer_dir_name="KDB생명",
        answer_json_name="운용지시서(KDB생명)_20260427.json",
        answer_csv_name="운용지시서(KDB생명)_20260427.csv",
        expected_orders=12,
        base_date="2026-04-27",
        kind="eml",
        strategy="single_prompt",
    ),
    CaseSpec(
        case_id="14_라이나_250826_xlsx",
        company="라이나생명",
        source_name="라이나_250826.xlsx",
        answer_dir_name="라이나생명",
        answer_json_name="라이나_250826.json",
        answer_csv_name="라이나_250826.csv",
        expected_orders=10,
        base_date="2025-08-26",
        kind="xlsx",
        strategy="single_prompt",
        sheet_name="당일",
        xlsx_range="A1:N12",
    ),
    CaseSpec(
        case_id="22_삼성자산_한화생명_설정_해지_내역서_20251128_eml",
        company="한화생명",
        source_name="삼성자산_한화생명_설정_해지_내역서_20251128.eml",
        answer_dir_name="한화생명",
        answer_json_name="삼성자산_한화생명_설정_해지_내역서_20251128.json",
        answer_csv_name="삼성자산_한화생명_설정_해지_내역서_20251128.csv",
        expected_orders=3,
        base_date="2025-11-28",
        kind="eml",
        strategy="single_prompt",
    ),
    CaseSpec(
        case_id="12_동양생명_20260318_html",
        company="동양생명",
        source_name="동양생명_20260318.html",
        answer_dir_name="동양생명",
        answer_json_name="동양생명_20260318.json",
        answer_csv_name="동양생명_20260318.csv",
        expected_orders=40,
        base_date="2026-03-18",
        kind="html",
        strategy="dongyang_two_stage",
        compact_fund_name=True,
    ),
    CaseSpec(
        case_id="37_카디프_251127_pdf",
        company="카디프생명",
        source_name="카디프_251127.pdf",
        answer_dir_name="카디프생명",
        answer_json_name="카디프_251127.json",
        answer_csv_name="카디프_251127.csv",
        expected_orders=46,
        base_date="2025-11-27",
        kind="pdf",
        strategy="single_prompt",
    ),
]


def normalize_name(value: str) -> str:
    return unicodedata.normalize("NFC", value)


def find_child(parent: Path, display_name: str) -> Path:
    exact = parent / display_name
    if exact.exists():
        return exact
    normalized_display = normalize_name(display_name)
    for child in parent.iterdir():
        if normalize_name(child.name) == normalized_display:
            return child
    raise FileNotFoundError(f"{display_name} not found under {parent}")


def resolve_case(spec: CaseSpec) -> dict[str, Any]:
    answer_dir = find_child(ANSWER_DIR, spec.answer_dir_name)
    return {
        "spec": spec,
        "case_id": spec.case_id,
        "company": spec.company,
        "source": find_child(DOCUMENT_DIR, spec.source_name),
        "answer_json": find_child(answer_dir, spec.answer_json_name),
        "answer_csv": find_child(answer_dir, spec.answer_csv_name),
        "expected_orders": spec.expected_orders,
        "base_date": spec.base_date,
        "kind": spec.kind,
        "strategy": spec.strategy,
        "compact_fund_name": spec.compact_fund_name,
        "sheet_name": spec.sheet_name,
        "xlsx_range": spec.xlsx_range,
    }


def load_env_file(path: Path) -> None:
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            continue
        if stripped.startswith("export "):
            stripped = stripped[7:].strip()
        key, value = stripped.split("=", 1)
        key = key.strip()
        value = value.strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
            value = value[1:-1]
        os.environ.setdefault(key, value)


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str), encoding="utf-8")


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def choose_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "/System/Library/Fonts/AppleSDGothicNeo.ttc",
        "/System/Library/Fonts/Supplemental/AppleGothic.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            try:
                return ImageFont.truetype(candidate, size=size, index=0)
            except Exception:
                continue
    return ImageFont.load_default()


def image_data_url(path: Path) -> str:
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:image/png;base64,{encoded}"


def chat_url(base_url: str) -> str:
    base = base_url.rstrip("/")
    if base.endswith("/v1"):
        return f"{base}/chat/completions"
    if base.endswith("/chat/completions"):
        return base
    return f"{base}/v1/chat/completions"


def excel_display(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return f"{value:,}"
    if isinstance(value, float):
        if abs(value - round(value)) < 1e-9:
            return f"{int(round(value)):,}"
        if abs(value) < 1:
            return f"{value:.6f}".rstrip("0").rstrip(".")
        return f"{value:,.6f}".rstrip("0").rstrip(".")
    return str(value)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.ImageFont, max_width: int, max_lines: int) -> list[str]:
    if not text:
        return [""]
    lines: list[str] = []
    for raw_line in str(text).splitlines() or [str(text)]:
        current = ""
        for char in raw_line:
            candidate = current + char
            bbox = draw.textbbox((0, 0), candidate, font=font)
            if bbox[2] - bbox[0] <= max_width or not current:
                current = candidate
            else:
                lines.append(current)
                current = char
        if current:
            lines.append(current)
    return lines[:max_lines] or [""]


def extract_eml_html(path: Path) -> str:
    message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    for part in message.walk():
        if part.get_content_type() == "text/html":
            return part.get_content()
    for part in message.walk():
        if part.get_content_type() == "text/plain":
            text = part.get_content()
            return f"<html><body><pre>{text}</pre></body></html>"
    raise RuntimeError(f"text/html or text/plain part not found: {path}")


def parse_html_table_for_rendering(html: str) -> tuple[list[RenderCell], int, int]:
    soup = BeautifulSoup(html, "lxml")
    table = soup.find("table")
    if table is None:
        raise RuntimeError("HTML table not found")
    cells: list[RenderCell] = []
    occupied: dict[tuple[int, int], bool] = {}
    max_col = 0
    rows = table.find_all("tr")
    for row_index, tr in enumerate(rows):
        col_index = 0
        for td in tr.find_all(["td", "th"], recursive=False):
            while occupied.get((row_index, col_index)):
                col_index += 1
            rowspan = int(td.get("rowspan") or 1)
            colspan = int(td.get("colspan") or 1)
            text = " ".join(td.get_text(" ", strip=True).split())
            cells.append(RenderCell(row=row_index, col=col_index, rowspan=rowspan, colspan=colspan, text=text))
            for dr in range(rowspan):
                for dc in range(colspan):
                    occupied[(row_index + dr, col_index + dc)] = True
            col_index += colspan
            max_col = max(max_col, col_index)
    return cells, len(rows), max_col


def render_html_table(path: Path, out_dir: Path, label: str) -> dict[str, Any]:
    html = extract_eml_html(path) if path.suffix.lower() == ".eml" else path.read_text(encoding="utf-8", errors="ignore")
    cells, row_count, col_count = parse_html_table_for_rendering(html)
    sample_text_by_col: dict[int, int] = {}
    for cell in cells:
        width = min(520, max(110, len(cell.text) * 12 + 36))
        per_col = max(110, width // max(1, cell.colspan))
        for offset in range(cell.colspan):
            sample_text_by_col[cell.col + offset] = max(sample_text_by_col.get(cell.col + offset, 110), per_col)
    col_widths = [sample_text_by_col.get(idx, 160) for idx in range(col_count)]
    if col_count <= 8:
        col_widths = [max(a, b) for a, b in zip((col_widths + [160] * 8)[:8], [420, 240, 150, 120, 230, 210, 250, 190])]
    col_widths = [min(520, max(120, width)) for width in col_widths]

    row_heights = [58] * row_count
    for cell in cells:
        if cell.row < 3:
            row_heights[cell.row] = max(row_heights[cell.row], 68)
        elif len(cell.text) > 45:
            row_heights[cell.row] = max(row_heights[cell.row], 72)

    scale = 2
    margin = 32
    width = sum(col_widths) + margin * 2
    height = sum(row_heights) + margin * 2
    image = Image.new("RGB", (width * scale, height * scale), "white")
    draw = ImageDraw.Draw(image)
    font = choose_font(18 * scale)
    header_font = choose_font(19 * scale)
    title_font = choose_font(26 * scale)

    x_positions = [margin * scale]
    for col_width in col_widths:
        x_positions.append(x_positions[-1] + col_width * scale)
    y_positions = [margin * scale]
    for row_height in row_heights:
        y_positions.append(y_positions[-1] + row_height * scale)

    for cell in cells:
        left = x_positions[cell.col]
        right = x_positions[min(cell.col + cell.colspan, col_count)]
        top = y_positions[cell.row]
        bottom = y_positions[min(cell.row + cell.rowspan, row_count)]
        fill = "#ddebf7" if cell.row < 2 else "#e2f0d9" if cell.row == 2 else "#fce4d6" if "소계" in cell.text or "총계" in cell.text else "#ffffff" if cell.row % 2 == 0 else "#f8fbff"
        draw.rectangle([left, top, right, bottom], fill=fill, outline="#777777", width=max(1, scale))
        if not cell.text:
            continue
        active_font = title_font if cell.row < 2 else header_font if cell.row == 2 else font
        pad = 8 * scale
        wrapped = wrap_text(draw, cell.text, active_font, max(20, (right - left) - pad * 2), max_lines=6)
        line_height = int(active_font.size * 1.15)
        text_y = top + max(pad, ((bottom - top) - line_height * len(wrapped)) // 2)
        for line in wrapped:
            draw.text((left + pad, text_y), line, fill="#111111", font=active_font)
            text_y += line_height

    out_dir.mkdir(parents=True, exist_ok=True)
    out_png = out_dir / f"{label}_table.png"
    image.save(out_png, optimize=True)
    return {
        "renderer": "html_table_pil",
        "source_path": str(path),
        "image_path": str(out_png),
        "html_table_count": html.lower().count("<table"),
        "html_row_count": row_count,
        "html_column_count": col_count,
        "image_width": image.width,
        "image_height": image.height,
        "image_sha256": sha256_file(out_png),
    }


def nonempty_bounds(worksheet: Any) -> tuple[int, int, int, int]:
    rows: list[int] = []
    cols: list[int] = []
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value not in (None, ""):
                rows.append(cell.row)
                cols.append(cell.column)
    if not rows:
        return 1, 1, 1, 1
    return min(rows), max(rows), min(cols), max(cols)


def render_xlsx_sheet(path: Path, out_dir: Path, sheet_name: str | None, xlsx_range: str | None = None) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
    if xlsx_range:
        min_col, min_row, max_col, max_row = range_boundaries(xlsx_range)
    else:
        min_row, max_row, min_col, max_col = nonempty_bounds(worksheet)
    rows = [
        [excel_display(worksheet.cell(row=row, column=col).value) for col in range(min_col, max_col + 1)]
        for row in range(min_row, max_row + 1)
    ]
    probe = Image.new("RGB", (10, 10), "white")
    probe_draw = ImageDraw.Draw(probe)
    font = choose_font(18)
    header_font = choose_font(18)
    col_widths: list[int] = []
    for col_index in range(max_col - min_col + 1):
        max_px = 110
        for row in rows:
            text = row[col_index]
            if not text:
                continue
            bbox = probe_draw.textbbox((0, 0), text, font=font)
            max_px = max(max_px, bbox[2] - bbox[0] + 28)
        col_widths.append(min(360, max(120, max_px)))
    row_heights = [54 if idx < 6 else 48 for idx in range(len(rows))]

    scale = 2
    margin = 28
    width = sum(col_widths) + margin * 2
    height = sum(row_heights) + margin * 2
    image = Image.new("RGB", (width * scale, height * scale), "white")
    draw = ImageDraw.Draw(image)
    font = choose_font(17 * scale)
    header_font = choose_font(18 * scale)
    title_font = choose_font(24 * scale)

    x_positions = [margin * scale]
    for col_width in col_widths:
        x_positions.append(x_positions[-1] + col_width * scale)
    y_positions = [margin * scale]
    for row_height in row_heights:
        y_positions.append(y_positions[-1] + row_height * scale)

    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            left = x_positions[c]
            right = x_positions[c + 1]
            top = y_positions[r]
            bottom = y_positions[r + 1]
            sheet_row = min_row + r
            fill = "#ddebf7" if sheet_row in {1, 3, 4, 6} else "#fce4d6" if "합계" in str(text) else "#ffffff" if r % 2 == 0 else "#f8fbff"
            draw.rectangle([left, top, right, bottom], fill=fill, outline="#8a8a8a", width=scale)
            if not text:
                continue
            active_font = title_font if sheet_row == 1 else header_font if sheet_row == 6 else font
            pad = 7 * scale
            wrapped = wrap_text(draw, text, active_font, max(20, (right - left) - pad * 2), max_lines=3)
            line_height = int(active_font.size * 1.18)
            text_y = top + max(pad, ((bottom - top) - line_height * len(wrapped)) // 2)
            for line in wrapped:
                draw.text((left + pad, text_y), line, fill="#111111", font=active_font)
                text_y += line_height

    out_dir.mkdir(parents=True, exist_ok=True)
    safe_sheet = re.sub(r"[^0-9A-Za-z가-힣_.-]+", "_", worksheet.title)
    out_png = out_dir / f"{path.stem}_{safe_sheet}_{min_row}_{max_row}.png"
    image.save(out_png, optimize=True)
    return {
        "renderer": "xlsx_openpyxl_pil",
        "source_path": str(path),
        "sheet_name": worksheet.title,
        "range": xlsx_range or f"R{min_row}C{min_col}:R{max_row}C{max_col}",
        "image_path": str(out_png),
        "image_width": image.width,
        "image_height": image.height,
        "image_sha256": sha256_file(out_png),
    }


def render_pdf_pages(path: Path, out_dir: Path, zoom: float = 1.7) -> dict[str, Any]:
    out_dir.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(str(path))
    pages: list[dict[str, Any]] = []
    for index, page in enumerate(doc, start=1):
        pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        out_png = out_dir / f"{path.stem}_page_{index}.png"
        pixmap.save(str(out_png))
        pages.append(
            {
                "page_index": index,
                "image_path": str(out_png),
                "image_width": pixmap.width,
                "image_height": pixmap.height,
                "image_sha256": sha256_file(out_png),
            }
        )
    return {"renderer": "pymupdf_page_render", "source_path": str(path), "page_count": len(pages), "pages": pages}


def render_html_to_png(html: str, out_png: Path, out_pdf: Path, zoom: float) -> dict[str, Any]:
    out_png.parent.mkdir(parents=True, exist_ok=True)
    HTML(string=html, base_url=str(ROOT)).write_pdf(str(out_pdf))
    doc = fitz.open(str(out_pdf))
    pages: list[Image.Image] = []
    for page in doc:
        pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        pages.append(Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples))
    width = max(page.width for page in pages)
    height = sum(page.height for page in pages)
    final = Image.new("RGB", (width, height), "white")
    y = 0
    for page in pages:
        final.paste(page, (0, y))
        y += page.height
    final.save(out_png)
    return {
        "image_path": str(out_png),
        "pdf_path": str(out_pdf),
        "page_count": len(pages),
        "image_width": final.width,
        "image_height": final.height,
        "image_sha256": sha256_file(out_png),
    }


def section_html(table: Any, title: str, page_width: int) -> str:
    return f"""<html><head><meta charset="utf-8"><style>
@page{{size:{page_width}px 1050px;margin:0}}
html,body{{margin:0;padding:0;background:white;color:#111;font-family:"Malgun Gothic","Apple SD Gothic Neo",Arial,sans-serif}}
.page{{width:{page_width - 40}px;padding:20px}}
h2{{font-size:28px;margin:0 0 12px;font-weight:900}}
table{{width:100%;border-collapse:collapse;table-layout:fixed;border:2px solid #333}}
th,td{{border:1.5px solid #666;padding:7px 10px;height:34px;font-size:20px;white-space:nowrap;overflow:visible;text-overflow:clip;line-height:1.15}}
th{{background:#eef1f5;text-align:center;font-weight:900}}
td.code{{font-weight:800}}
td.num,th.num{{text-align:right;font-variant-numeric:tabular-nums}}
.totals td{{font-weight:800;background:#f7f7f7}}
col:nth-child(1){{width:135px!important}}
col:nth-child(2){{width:650px!important}}
col:nth-child(3),col:nth-child(4),col:nth-child(5),col:nth-child(6){{width:300px!important}}
</style></head><body><div class="page"><h2>{title}</h2>{str(table)}</div></body></html>"""


def row_meta(table: Any) -> dict[str, int]:
    rows = table.find_all("tr")
    body = table.find("tbody") or table
    data_rows = [tr for tr in body.find_all("tr", recursive=False) if "totals" not in (tr.get("class") or [])]
    return {
        "rows": len(rows),
        "data_rows": len(data_rows),
        "max_columns": max((len(row.find_all(["td", "th"])) for row in rows), default=0),
    }


def render_dongyang_sections(path: Path, out_dir: Path) -> dict[str, Any]:
    html = path.read_text(encoding="utf-8", errors="ignore")
    soup = BeautifulSoup(html, "html.parser")
    tables = soup.find_all("table")
    if len(tables) < 2:
        raise RuntimeError(f"{path.name} requires at least two tables")
    settlement_meta = row_meta(tables[0])
    forecast_meta = row_meta(tables[1])
    settlement_meta.update(
        render_html_to_png(
            section_html(tables[0], "정산내역", 1800),
            out_dir / "settlement_table.png",
            out_dir / "settlement_table.pdf",
            zoom=1.5,
        )
    )
    forecast_meta.update(
        render_html_to_png(
            section_html(tables[1], "예상내역", 2100),
            out_dir / "forecast_table.png",
            out_dir / "forecast_table.pdf",
            zoom=1.5,
        )
    )
    return {
        "renderer": "weasyprint_wide_page_section_render",
        "source_path": str(path),
        "settlement_table": settlement_meta,
        "forecast_table": forecast_meta,
    }


def render_assets(case: dict[str, Any], asset_root: Path) -> dict[str, Any]:
    image_dir = asset_root / case["case_id"] / "input_image"
    image_dir.mkdir(parents=True, exist_ok=True)
    source = case["source"]
    if case["strategy"] == "dongyang_two_stage":
        manifest = render_dongyang_sections(source, image_dir)
    elif case["kind"] == "xlsx":
        manifest = render_xlsx_sheet(source, image_dir, case.get("sheet_name"), case.get("xlsx_range"))
    elif case["kind"] == "pdf":
        manifest = render_pdf_pages(source, image_dir)
    elif case["kind"] == "eml":
        manifest = render_html_table(source, image_dir, case["case_id"])
    else:
        raise RuntimeError(f"unsupported render kind: {case['kind']}")
    write_json(asset_root / case["case_id"] / "image_manifest.json", manifest)
    return manifest


def base_output_shape(case: dict[str, Any], model: str) -> str:
    return json.dumps(
        {
            "file_name": case["source"].name,
            "source_path": str(case["source"]),
            "model_name": model,
            "base_date": "YYYY-MM-DD",
            "status": "COMPLETED",
            "reason": None,
            "issues": [],
            "orders": [
                {
                    "fund_code": "string",
                    "fund_name": "string",
                    "settle_class": "1 or 2",
                    "order_type": "3 or 1",
                    "base_date": "YYYY-MM-DD",
                    "t_day": "01",
                    "transfer_amount": "string",
                }
            ],
        },
        ensure_ascii=False,
        indent=2,
    )


def render_template(path: Path, values: dict[str, Any]) -> str:
    template = path.read_text(encoding="utf-8")
    rendered = template
    for key, value in values.items():
        rendered = rendered.replace(f"{{{{{key}}}}}", str(value))
    unresolved = sorted(set(re.findall(r"\{\{[A-Z0-9_]+\}\}", rendered)))
    if unresolved:
        raise RuntimeError(f"unresolved placeholders in {path}: {', '.join(unresolved)}")
    return rendered


def template_values(case: dict[str, Any], model: str, **extra: Any) -> dict[str, Any]:
    values: dict[str, Any] = {
        "BASE_DATE": case["base_date"],
        "MODEL": model,
        "SOURCE_NAME": case["source"].name,
        "OUTPUT_SHAPE": base_output_shape(case, model),
    }
    values.update(extra)
    return values


def single_case_prompt(case: dict[str, Any], model: str) -> str:
    path = PROMPT_DIR / f"{case['case_id']}.promt"
    return render_template(path, template_values(case, model))


def stage_prompt(case: dict[str, Any], model: str, stage: str, **extra: Any) -> str:
    path = PROMPT_DIR / f"{case['case_id']}__{stage}.promt"
    return render_template(path, template_values(case, model, **extra))


def profile_payload(profile: Profile) -> dict[str, Any]:
    return {
        "enable_thinking": profile.enable_thinking,
        "separate_reasoning": True,
        "temperature": profile.temperature,
        "top_p": profile.top_p,
        "top_k": profile.top_k,
        "min_p": profile.min_p,
        "presence_penalty": profile.presence_penalty,
        "repetition_penalty": profile.repetition_penalty,
        "max_tokens": profile.max_tokens,
        "stream": False,
        "response_format": {"type": "json_object"},
    }


def call_novita_json(
    prompt: str,
    images: list[Path],
    profile: Profile,
    label: str,
    response_dir: Path,
    *,
    endpoint: str,
    api_key: str,
    model: str,
    timeout: int,
) -> tuple[dict[str, Any] | None, dict[str, Any], dict[str, Any]]:
    response_dir.mkdir(parents=True, exist_ok=True)
    content: list[dict[str, Any]] = [{"type": "text", "text": prompt}]
    content.extend({"type": "image_url", "image_url": {"url": image_data_url(image)}} for image in images)
    body: dict[str, Any] = {
        "model": model,
        "messages": [{"role": "user", "content": content}],
        **profile_payload(profile),
    }
    safe_request = {
        "url": endpoint,
        "model": model,
        "message_count": 1,
        "content_parts": ["text", *["image_url" for _ in images]],
        **profile_payload(profile),
    }
    started = time.perf_counter()
    try:
        response = requests.post(
            endpoint,
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            json=body,
            timeout=timeout,
        )
        elapsed = round(time.perf_counter() - started, 3)
    except requests.Timeout as exc:
        elapsed = round(time.perf_counter() - started, 3)
        raise RunBlockedError(f"{label} timeout after {elapsed}s: {exc}") from exc
    except requests.RequestException as exc:
        elapsed = round(time.perf_counter() - started, 3)
        raise RunBlockedError(f"{label} request error after {elapsed}s: {exc}") from exc

    try:
        response_json: dict[str, Any] | None = response.json()
    except Exception:
        response_json = None

    raw = {
        "request": safe_request,
        "http_status": response.status_code,
        "elapsed_seconds": elapsed,
        "response_json": response_json,
        "response_text": response.text if response_json is None else None,
    }
    write_json(response_dir / f"{label}_raw_response.json", raw)

    if response.status_code >= 400:
        raise RunBlockedError(f"{label} HTTP {response.status_code}: {response.text[:1000]}")

    choices = (response_json or {}).get("choices") or []
    if not choices:
        parts = {"label": label, "content": None, "finish_reason": None, "usage": (response_json or {}).get("usage"), "elapsed_seconds": elapsed}
        write_json(response_dir / f"{label}_message_parts.json", parts)
        return None, {"parse_ok": False, "error": "choices is empty"}, parts

    choice = choices[0]
    message = choice.get("message") or {}
    content_text = message.get("content")
    parts = {
        "label": label,
        "finish_reason": choice.get("finish_reason"),
        "elapsed_seconds": elapsed,
        "usage": (response_json or {}).get("usage"),
        "reasoning_content_present": message.get("reasoning_content") is not None or message.get("reasoning") is not None,
        "reasoning_content": message.get("reasoning_content", message.get("reasoning")),
        "content": content_text,
    }
    write_json(response_dir / f"{label}_message_parts.json", parts)
    if not isinstance(content_text, str):
        return None, {"parse_ok": False, "error": f"message.content is not string: {type(content_text).__name__}"}, parts
    parsed, parse_info = parse_json_content(content_text)
    return parsed, parse_info, parts


def parse_json_content(content: str) -> tuple[dict[str, Any] | None, dict[str, Any]]:
    stripped = (content or "").strip()
    attempts: list[tuple[str, str]] = [("raw", stripped)]
    fence_match = re.fullmatch(r"```(?:json)?\s*(.*?)\s*```", stripped, flags=re.IGNORECASE | re.DOTALL)
    if fence_match:
        attempts.append(("markdown_fence", fence_match.group(1).strip()))
    first = stripped.find("{")
    last = stripped.rfind("}")
    if first != -1 and last > first:
        attempts.append(("object_slice", stripped[first : last + 1]))

    errors: list[dict[str, str]] = []
    for mode, candidate in attempts:
        try:
            parsed = json.loads(candidate)
        except Exception as exc:
            errors.append({"mode": mode, "error": str(exc)})
            continue
        if not isinstance(parsed, dict):
            return None, {"parse_ok": False, "error": "parsed content is not JSON object", "mode": mode}
        return parsed, {"parse_ok": True, "normalized_mode": mode, "content_preview": stripped[:1000]}
    return None, {"parse_ok": False, "errors": errors, "content_preview": stripped[:3000]}


def clean_amount(value: Any) -> str:
    return str(value or "").strip().replace(" ", "")


def amount_nonzero(value: Any) -> bool:
    text = clean_amount(value).replace(",", "").replace("-", "").replace("−", "")
    if text in {"", ".", "0", "0.0"}:
        return False
    digits = re.sub(r"[^0-9.]", "", text)
    if not digits:
        return False
    try:
        return float(digits) != 0
    except Exception:
        return bool(re.search(r"[1-9]", digits))


def output_amount(value: Any) -> str:
    return clean_amount(value).replace("-", "").replace("−", "")


def rows_by_key(rows: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    for row in rows:
        code = str(row.get("fund_code") or "").strip()
        if code and code not in result:
            result[code] = row
    return result


def assemble_dongyang_payload(
    case: dict[str, Any],
    model: str,
    settlement: dict[str, Any] | None,
    forecast: dict[str, Any] | None,
) -> tuple[dict[str, Any], dict[str, Any]]:
    orders: list[dict[str, str]] = []
    trace: list[dict[str, Any]] = []
    base_date = case["base_date"]
    settlement_rows = (settlement or {}).get("rows") or []
    forecast_rows = (forecast or {}).get("rows") or []
    settlement_name_by_code = {code: str(row.get("fund_name") or "").strip() for code, row in rows_by_key(settlement_rows).items()}

    for row in settlement_rows:
        code = str(row.get("fund_code") or "").strip()
        name = str(row.get("fund_name") or "").strip()
        if not code or "합계" in code or "합계" in name:
            continue
        if amount_nonzero(row.get("setting_amount")):
            order = {"fund_code": code, "fund_name": name, "settle_class": "2", "order_type": "3", "base_date": base_date, "t_day": "01", "transfer_amount": output_amount(row.get("setting_amount"))}
            orders.append(order)
            trace.append({"source": "settlement.setting_amount", "row": row, "order": order})
        if amount_nonzero(row.get("redemption_amount")):
            order = {"fund_code": code, "fund_name": name, "settle_class": "2", "order_type": "1", "base_date": base_date, "t_day": "01", "transfer_amount": output_amount(row.get("redemption_amount"))}
            orders.append(order)
            trace.append({"source": "settlement.redemption_amount", "row": row, "order": order})

    for row in forecast_rows:
        code = str(row.get("fund_code") or "").strip()
        name = settlement_name_by_code.get(code, str(row.get("fund_name") or "").strip())
        if not code or "합계" in code or "합계" in name:
            continue
        if amount_nonzero(row.get("date2_setting")):
            order = {"fund_code": code, "fund_name": name, "settle_class": "1", "order_type": "3", "base_date": base_date, "t_day": "03", "transfer_amount": output_amount(row.get("date2_setting"))}
            orders.append(order)
            trace.append({"source": "forecast.date2_setting", "row": row, "order": order})
        if amount_nonzero(row.get("date2_redemption")):
            order = {"fund_code": code, "fund_name": name, "settle_class": "1", "order_type": "1", "base_date": base_date, "t_day": "03", "transfer_amount": output_amount(row.get("date2_redemption"))}
            orders.append(order)
            trace.append({"source": "forecast.date2_redemption", "row": row, "order": order})

    issues: list[str] = []
    if settlement and settlement.get("issues"):
        issues.extend(f"settlement:{item}" for item in settlement.get("issues") or [])
    if forecast and forecast.get("issues"):
        issues.extend(f"forecast:{item}" for item in forecast.get("issues") or [])
    payload = {
        "file_name": case["source"].name,
        "source_path": str(case["source"]),
        "model_name": model,
        "base_date": base_date,
        "status": "COMPLETED" if not issues else "FAILED",
        "reason": None if not issues else "LLM extraction reported issues",
        "issues": issues,
        "orders": orders,
    }
    return payload, {"orders": trace, "settlement_name_by_code": settlement_name_by_code}


def normalize_amount(value: Any) -> str:
    text = str(value or "").strip().replace(",", "").replace("-", "").replace("−", "")
    return text[:-2] if text.endswith(".0") else text


def normalize_scalar(value: Any) -> str:
    return "" if value is None else str(value).strip()


def normalize_fund_name(value: Any, compact: bool) -> str:
    text = normalize_scalar(value)
    if compact:
        return re.sub(r"\s+", "", text)
    return " ".join(text.split())


def canonical_order(order: dict[str, Any], *, compact_fund_name: bool) -> dict[str, str]:
    settle = normalize_scalar(order.get("settle_class"))
    order_type = normalize_scalar(order.get("order_type"))
    t_day = normalize_scalar(order.get("t_day"))
    fund_code = order.get("fund_code", order.get("\ufefffund_code", ""))
    return {
        "fund_code": normalize_scalar(fund_code),
        "fund_name": normalize_fund_name(order.get("fund_name"), compact_fund_name),
        "base_date": normalize_scalar(order.get("base_date")),
        "t_day": t_day.zfill(2) if t_day else "",
        "transfer_amount": normalize_amount(order.get("transfer_amount")),
        "settle_class": SETTLE_TO_CODE.get(settle, settle),
        "order_type": ORDER_TO_CODE.get(order_type, order_type),
    }


def canonical_orders(rows: list[dict[str, Any]], *, compact_fund_name: bool) -> list[dict[str, str]]:
    return sorted(
        [canonical_order(row, compact_fund_name=compact_fund_name) for row in rows],
        key=lambda row: tuple(row.get(field, "") for field in CSV_FIELDS),
    )


def core_payload(payload: dict[str, Any], *, compact_fund_name: bool) -> dict[str, Any]:
    return {
        "status": payload.get("status"),
        "base_date": payload.get("base_date"),
        "reason": payload.get("reason"),
        "issues": payload.get("issues") or [],
        "orders": canonical_orders(payload.get("orders") or [], compact_fund_name=compact_fund_name),
    }


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return [dict(row) for row in csv.DictReader(handle)]


def write_result_csv(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=CSV_FIELDS)
        writer.writeheader()
        for order in payload.get("orders") or []:
            writer.writerow({field: order.get(field, "") for field in CSV_FIELDS})


def order_multiset(rows: list[dict[str, Any]], *, compact_fund_name: bool) -> Counter[str]:
    return Counter(json.dumps(row, ensure_ascii=False, sort_keys=True) for row in canonical_orders(rows, compact_fund_name=compact_fund_name))


def diff_orders(result_orders: list[dict[str, Any]], answer_orders: list[dict[str, Any]], *, compact_fund_name: bool) -> dict[str, Any]:
    result_set = order_multiset(result_orders, compact_fund_name=compact_fund_name)
    answer_set = order_multiset(answer_orders, compact_fund_name=compact_fund_name)
    missing: list[dict[str, Any]] = []
    extra: list[dict[str, Any]] = []
    for key, count in answer_set.items():
        missing.extend([json.loads(key)] * max(0, count - result_set.get(key, 0)))
    for key, count in result_set.items():
        extra.extend([json.loads(key)] * max(0, count - answer_set.get(key, 0)))
    return {"missing_from_result": missing, "extra_in_result": extra}


def validate_payload(case: dict[str, Any], payload: dict[str, Any], result_csv: Path) -> dict[str, Any]:
    answer_payload = read_json(case["answer_json"])
    compact = bool(case.get("compact_fund_name"))
    result_csv_rows = read_csv_rows(result_csv)
    answer_csv_rows = read_csv_rows(case["answer_csv"])
    checks = {
        "status_completed": payload.get("status") == "COMPLETED",
        "base_date_expected": payload.get("base_date") == case["base_date"],
        "orders_expected": len(payload.get("orders") or []) == case["expected_orders"],
        "issues_empty": (payload.get("issues") or []) == [],
        "json_core_exact": core_payload(payload, compact_fund_name=compact) == core_payload(answer_payload, compact_fund_name=compact),
        "csv_canonical_exact": order_multiset(result_csv_rows, compact_fund_name=compact) == order_multiset(answer_csv_rows, compact_fund_name=compact),
    }
    return {
        "case_id": case["case_id"],
        "verdict": "PASS" if all(checks.values()) else "FAIL",
        "checks": checks,
        "status": payload.get("status"),
        "base_date": payload.get("base_date"),
        "orders": len(payload.get("orders") or []),
        "expected_orders": case["expected_orders"],
        "diff": diff_orders(payload.get("orders") or [], answer_payload.get("orders") or [], compact_fund_name=compact),
    }


def summarize_parts(parts: list[dict[str, Any]]) -> dict[str, Any]:
    elapsed = round(sum(float(part.get("elapsed_seconds") or 0) for part in parts), 3)
    prompt_tokens = 0
    completion_tokens = 0
    total_tokens = 0
    usage_by_call: dict[str, Any] = {}
    reasoning_present = False
    finish_reasons: dict[str, Any] = {}
    for part in parts:
        label = part.get("label") or f"call_{len(usage_by_call) + 1}"
        usage = part.get("usage") or {}
        usage_by_call[label] = usage
        prompt_tokens += int(usage.get("prompt_tokens") or usage.get("input_tokens") or 0)
        completion_tokens += int(usage.get("completion_tokens") or usage.get("output_tokens") or 0)
        total_tokens += int(usage.get("total_tokens") or 0)
        reasoning_present = reasoning_present or bool(part.get("reasoning_content_present"))
        finish_reasons[label] = part.get("finish_reason")
    if not total_tokens:
        total_tokens = prompt_tokens + completion_tokens
    return {
        "llm_call_count": len(parts),
        "elapsed_seconds": elapsed,
        "usage": {
            "prompt_tokens": prompt_tokens,
            "completion_tokens": completion_tokens,
            "total_tokens": total_tokens,
            "by_call": usage_by_call,
        },
        "finish_reasons": finish_reasons,
        "reasoning_content_present": reasoning_present,
    }


def blocked_summary(
    case: dict[str, Any],
    cycle: int,
    profile: Profile,
    execution_dir: Path,
    reason: str,
    parts: list[dict[str, Any]] | None = None,
) -> dict[str, Any]:
    call_summary = summarize_parts(parts or [])
    summary = {
        "case_id": case["case_id"],
        "company": case["company"],
        "cycle": cycle,
        "profile": profile.name,
        "verdict": "BLOCKED",
        "blocked_reason": reason,
        "status": None,
        "base_date": None,
        "orders": 0,
        "expected_orders": case["expected_orders"],
        "json_core_exact": False,
        "csv_canonical_exact": False,
        **call_summary,
        "paths": {"execution_dir": str(execution_dir)},
    }
    write_json(execution_dir / "case_summary.json", summary)
    return summary


def fail_summary(
    case: dict[str, Any],
    cycle: int,
    profile: Profile,
    execution_dir: Path,
    reason: str,
    parts: list[dict[str, Any]],
    parse_info: dict[str, Any] | None = None,
) -> dict[str, Any]:
    call_summary = summarize_parts(parts)
    validation = {
        "case_id": case["case_id"],
        "verdict": "FAIL",
        "reason": reason,
        "parse_info": parse_info,
        "checks": {
            "status_completed": False,
            "base_date_expected": False,
            "orders_expected": False,
            "issues_empty": False,
            "json_core_exact": False,
            "csv_canonical_exact": False,
        },
    }
    write_json(execution_dir / "validation_summary.json", validation)
    summary = {
        "case_id": case["case_id"],
        "company": case["company"],
        "cycle": cycle,
        "profile": profile.name,
        "verdict": "FAIL",
        "failure_reason": reason,
        "status": None,
        "base_date": None,
        "orders": 0,
        "expected_orders": case["expected_orders"],
        "json_core_exact": False,
        "csv_canonical_exact": False,
        **call_summary,
        "paths": {"execution_dir": str(execution_dir), "validation_summary": str(execution_dir / "validation_summary.json")},
    }
    write_json(execution_dir / "case_summary.json", summary)
    return summary


def images_for_single_case(case: dict[str, Any], manifest: dict[str, Any]) -> list[Path]:
    if case["kind"] == "pdf":
        return [Path(page["image_path"]) for page in manifest["pages"]]
    return [Path(manifest["image_path"])]


def run_single_prompt_case(
    case: dict[str, Any],
    manifest: dict[str, Any],
    cycle: int,
    profile: Profile,
    execution_dir: Path,
    *,
    endpoint: str,
    api_key: str,
    model: str,
    timeout: int,
) -> dict[str, Any]:
    prompt = single_case_prompt(case, model)
    execution_dir.mkdir(parents=True, exist_ok=True)
    (execution_dir / "prompt.md").write_text(prompt, encoding="utf-8")
    response_dir = execution_dir / "response"
    try:
        parsed, parse_info, parts = call_novita_json(
            prompt,
            images_for_single_case(case, manifest),
            profile,
            "extract",
            response_dir,
            endpoint=endpoint,
            api_key=api_key,
            model=model,
            timeout=timeout,
        )
    except RunBlockedError as exc:
        return blocked_summary(case, cycle, profile, execution_dir, str(exc))

    if parsed is None:
        return fail_summary(case, cycle, profile, execution_dir, "JSON parse failed", [parts], parse_info)

    result_json = execution_dir / "result" / f"{case['case_id']}.json"
    result_csv = execution_dir / "result" / f"{case['case_id']}.csv"
    write_json(result_json, parsed)
    write_result_csv(result_csv, parsed)
    validation = validate_payload(case, parsed, result_csv)
    write_json(execution_dir / "validation_summary.json", validation)
    call_summary = summarize_parts([parts])
    summary = {
        "case_id": case["case_id"],
        "company": case["company"],
        "cycle": cycle,
        "profile": profile.name,
        "verdict": validation["verdict"],
        "status": parsed.get("status"),
        "base_date": parsed.get("base_date"),
        "orders": len(parsed.get("orders") or []),
        "expected_orders": case["expected_orders"],
        "json_core_exact": validation["checks"]["json_core_exact"],
        "csv_canonical_exact": validation["checks"]["csv_canonical_exact"],
        **call_summary,
        "paths": {
            "execution_dir": str(execution_dir),
            "result_json": str(result_json),
            "result_csv": str(result_csv),
            "validation_summary": str(execution_dir / "validation_summary.json"),
            "answer_json": str(case["answer_json"]),
            "answer_csv": str(case["answer_csv"]),
        },
    }
    write_json(execution_dir / "case_summary.json", summary)
    return summary


def run_dongyang_case(
    case: dict[str, Any],
    manifest: dict[str, Any],
    cycle: int,
    profile: Profile,
    execution_dir: Path,
    *,
    endpoint: str,
    api_key: str,
    model: str,
    timeout: int,
) -> dict[str, Any]:
    execution_dir.mkdir(parents=True, exist_ok=True)
    response_dir = execution_dir / "response"
    intermediate_dir = execution_dir / "intermediate"
    settlement_image = Path(manifest["settlement_table"]["image_path"])
    forecast_image = Path(manifest["forecast_table"]["image_path"])
    settlement_prompt = stage_prompt(case, model, "settlement", ROW_COUNT=manifest["settlement_table"]["data_rows"])
    (execution_dir / "prompt_settlement.md").write_text(settlement_prompt, encoding="utf-8")
    parts: list[dict[str, Any]] = []
    try:
        settlement, settlement_parse, settlement_parts = call_novita_json(
            settlement_prompt,
            [settlement_image],
            profile,
            "settlement",
            response_dir,
            endpoint=endpoint,
            api_key=api_key,
            model=model,
            timeout=timeout,
        )
    except RunBlockedError as exc:
        return blocked_summary(case, cycle, profile, execution_dir, str(exc), parts)
    parts.append(settlement_parts)
    if settlement is None:
        return fail_summary(case, cycle, profile, execution_dir, "settlement JSON parse failed", parts, settlement_parse)
    write_json(intermediate_dir / "settlement_rows.json", settlement)

    settlement_compact = json.dumps(settlement or {}, ensure_ascii=False, separators=(",", ":"))
    forecast_prompt = stage_prompt(
        case,
        model,
        "forecast",
        ROW_COUNT=manifest["forecast_table"]["data_rows"],
        SETTLEMENT_ROWS_JSON=settlement_compact,
    )
    (execution_dir / "prompt_forecast.md").write_text(forecast_prompt, encoding="utf-8")
    try:
        forecast, forecast_parse, forecast_parts = call_novita_json(
            forecast_prompt,
            [forecast_image],
            profile,
            "forecast",
            response_dir,
            endpoint=endpoint,
            api_key=api_key,
            model=model,
            timeout=timeout,
        )
    except RunBlockedError as exc:
        return blocked_summary(case, cycle, profile, execution_dir, str(exc), parts)
    parts.append(forecast_parts)
    if forecast is None:
        return fail_summary(case, cycle, profile, execution_dir, "forecast JSON parse failed", parts, forecast_parse)
    write_json(intermediate_dir / "forecast_rows.json", forecast)

    payload, assembly_trace = assemble_dongyang_payload(case, model, settlement, forecast)
    write_json(intermediate_dir / "assembly_trace.json", assembly_trace)
    result_json = execution_dir / "result" / f"{case['case_id']}.json"
    result_csv = execution_dir / "result" / f"{case['case_id']}.csv"
    write_json(result_json, payload)
    write_result_csv(result_csv, payload)
    validation = validate_payload(case, payload, result_csv)
    write_json(execution_dir / "validation_summary.json", validation)
    call_summary = summarize_parts(parts)
    summary = {
        "case_id": case["case_id"],
        "company": case["company"],
        "cycle": cycle,
        "profile": profile.name,
        "verdict": validation["verdict"],
        "status": payload.get("status"),
        "base_date": payload.get("base_date"),
        "orders": len(payload.get("orders") or []),
        "expected_orders": case["expected_orders"],
        "json_core_exact": validation["checks"]["json_core_exact"],
        "csv_canonical_exact": validation["checks"]["csv_canonical_exact"],
        **call_summary,
        "paths": {
            "execution_dir": str(execution_dir),
            "result_json": str(result_json),
            "result_csv": str(result_csv),
            "validation_summary": str(execution_dir / "validation_summary.json"),
            "answer_json": str(case["answer_json"]),
            "answer_csv": str(case["answer_csv"]),
        },
    }
    write_json(execution_dir / "case_summary.json", summary)
    return summary


def run_case_execution(
    case: dict[str, Any],
    manifest: dict[str, Any],
    cycle: int,
    profile: Profile,
    execution_dir: Path,
    *,
    endpoint: str,
    api_key: str,
    model: str,
    timeout: int,
) -> dict[str, Any]:
    try:
        if case["strategy"] == "dongyang_two_stage":
            return run_dongyang_case(case, manifest, cycle, profile, execution_dir, endpoint=endpoint, api_key=api_key, model=model, timeout=timeout)
        return run_single_prompt_case(case, manifest, cycle, profile, execution_dir, endpoint=endpoint, api_key=api_key, model=model, timeout=timeout)
    except Exception as exc:
        trace_path = execution_dir / "unexpected_error.txt"
        trace_path.parent.mkdir(parents=True, exist_ok=True)
        trace_path.write_text(traceback.format_exc(), encoding="utf-8")
        return fail_summary(case, cycle, profile, execution_dir, f"unexpected error: {exc}", [])


def aggregate_profile_summary(summaries: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    for profile in sorted({summary["profile"] for summary in summaries}):
        rows = [summary for summary in summaries if summary["profile"] == profile]
        completed = [row for row in rows if row["verdict"] != "BLOCKED"]
        result.append(
            {
                "profile": profile,
                "total": len(rows),
                "pass": sum(1 for row in rows if row["verdict"] == "PASS"),
                "fail": sum(1 for row in rows if row["verdict"] == "FAIL"),
                "blocked": sum(1 for row in rows if row["verdict"] == "BLOCKED"),
                "avg_elapsed_seconds": round(sum(float(row.get("elapsed_seconds") or 0) for row in completed) / len(completed), 3) if completed else 0,
                "avg_total_tokens": round(sum(int((row.get("usage") or {}).get("total_tokens") or 0) for row in completed) / len(completed), 1) if completed else 0,
            }
        )
    return result


def relative(path: str | Path) -> str:
    try:
        return str(Path(path).relative_to(ROOT))
    except Exception:
        return str(path)


def write_report(
    report_path: Path,
    *,
    run_id: str,
    artifact_root: Path,
    cases: list[dict[str, Any]],
    profiles: list[Profile],
    cycles: int,
    endpoint: str,
    model: str,
    summaries: list[dict[str, Any]],
) -> None:
    profile_summary = aggregate_profile_summary(summaries)
    verdict_counts = Counter(summary["verdict"] for summary in summaries)
    overall = "PASS" if verdict_counts.get("FAIL", 0) == 0 and verdict_counts.get("BLOCKED", 0) == 0 else "REVIEW"
    lines = [
        "# Novita Qwen3.5 지시서 추출 파라미터 테스트 보고서",
        "",
        f"- 실행 일시: `{datetime.now(KST).strftime('%Y-%m-%d %H:%M:%S KST')}`",
        f"- run id: `{run_id}`",
        f"- overall verdict: `{overall}`",
        f"- model: `{model}`",
        f"- endpoint: `{endpoint}`",
        f"- cycles: `{cycles}`",
        f"- artifact root: [{artifact_root.name}]({artifact_root})",
        "- deterministic 보강 사용 여부: `false`",
        "- API key 저장 여부: `false`",
        "",
        "## 테스트 대상",
        "",
        "| Case | 거래처 | Source | Expected | Base date | Strategy |",
        "| --- | --- | --- | ---: | --- | --- |",
    ]
    for case in cases:
        lines.append(
            f"| `{case['case_id']}` | {case['company']} | [{case['source'].name}]({case['source']}) | "
            f"`{case['expected_orders']}` | `{case['base_date']}` | `{case['strategy']}` |"
        )
    lines.extend(["", "## 파라미터 프로필", "", "| Profile | enable_thinking | temperature | top_p | top_k | min_p | presence_penalty | repetition_penalty | max_tokens |", "| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: |"])
    for profile in profiles:
        lines.append(
            f"| `{profile.name}` | `{str(profile.enable_thinking).lower()}` | `{profile.temperature}` | `{profile.top_p}` | "
            f"`{profile.top_k}` | `{profile.min_p}` | `{profile.presence_penalty}` | `{profile.repetition_penalty}` | `{profile.max_tokens}` |"
        )
    lines.extend(["", "## Verdict Summary", "", f"- PASS: `{verdict_counts.get('PASS', 0)}`", f"- FAIL: `{verdict_counts.get('FAIL', 0)}`", f"- BLOCKED: `{verdict_counts.get('BLOCKED', 0)}`", ""])
    lines.extend(["| Profile | Total | PASS | FAIL | BLOCKED | Avg elapsed(s) | Avg total tokens |", "| --- | ---: | ---: | ---: | ---: | ---: | ---: |"])
    for row in profile_summary:
        lines.append(
            f"| `{row['profile']}` | `{row['total']}` | `{row['pass']}` | `{row['fail']}` | `{row['blocked']}` | "
            f"`{row['avg_elapsed_seconds']}` | `{row['avg_total_tokens']}` |"
        )
    lines.extend(["", "## 실행 결과", "", "| Cycle | Profile | Case | Verdict | Orders | Expected | JSON exact | CSV exact | Calls | Elapsed(s) | Tokens | Reasoning | Artifact |", "| ---: | --- | --- | --- | ---: | ---: | --- | --- | ---: | ---: | ---: | --- | --- |"])
    for summary in summaries:
        usage = summary.get("usage") or {}
        artifact = summary.get("paths", {}).get("execution_dir", "")
        lines.append(
            f"| `{summary['cycle']}` | `{summary['profile']}` | `{summary['case_id']}` | `{summary['verdict']}` | "
            f"`{summary.get('orders', 0)}` | `{summary.get('expected_orders', 0)}` | `{summary.get('json_core_exact')}` | "
            f"`{summary.get('csv_canonical_exact')}` | `{summary.get('llm_call_count', 0)}` | `{summary.get('elapsed_seconds', 0)}` | "
            f"`{usage.get('total_tokens', 0)}` | `{summary.get('reasoning_content_present')}` | "
            f"[artifact]({artifact}) |"
        )
    lines.extend(
        [
            "",
            "## LLM 처리 시간 및 Token Usage",
            "",
            "| Cycle | Profile | Case | Calls | LLM elapsed(s) | Prompt tokens | Completion tokens | Total tokens | Finish reasons |",
            "| ---: | --- | --- | ---: | ---: | ---: | ---: | ---: | --- |",
        ]
    )
    for summary in summaries:
        usage = summary.get("usage") or {}
        finish_reasons = summary.get("finish_reasons") or {}
        lines.append(
            f"| `{summary['cycle']}` | `{summary['profile']}` | `{summary['case_id']}` | `{summary.get('llm_call_count', 0)}` | "
            f"`{summary.get('elapsed_seconds', 0)}` | `{usage.get('prompt_tokens', 0)}` | "
            f"`{usage.get('completion_tokens', 0)}` | `{usage.get('total_tokens', 0)}` | "
            f"`{json.dumps(finish_reasons, ensure_ascii=False)}` |"
        )
    multi_call_summaries = [summary for summary in summaries if summary.get("llm_call_count", 0) > 1]
    if multi_call_summaries:
        lines.extend(["", "### 호출별 Usage", ""])
        for summary in multi_call_summaries:
            usage_by_call = (summary.get("usage") or {}).get("by_call") or {}
            lines.append(f"#### `{summary['cycle']}` / `{summary['profile']}` / `{summary['case_id']}`")
            lines.append("")
            lines.append("| Call | Prompt tokens | Completion tokens | Total tokens |")
            lines.append("| --- | ---: | ---: | ---: |")
            for label, usage in usage_by_call.items():
                lines.append(
                    f"| `{label}` | `{usage.get('prompt_tokens', usage.get('input_tokens', 0))}` | "
                    f"`{usage.get('completion_tokens', usage.get('output_tokens', 0))}` | `{usage.get('total_tokens', 0)}` |"
                )
            lines.append("")
    failures = [summary for summary in summaries if summary["verdict"] != "PASS"]
    if failures:
        lines.extend(["", "## 검토 필요 항목", ""])
        for summary in failures:
            validation_path = Path(summary.get("paths", {}).get("validation_summary", summary.get("paths", {}).get("execution_dir", "")) or "")
            reason = summary.get("blocked_reason") or summary.get("failure_reason") or "validation mismatch"
            lines.append(f"### `{summary['cycle']}` / `{summary['profile']}` / `{summary['case_id']}`")
            lines.append(f"- verdict: `{summary['verdict']}`")
            lines.append(f"- reason: `{reason}`")
            lines.append("- analysis basis: `prompt-only result; no deterministic repair was applied`")
            if validation_path.is_file():
                lines.append(f"- validation: [validation_summary.json]({validation_path})")
                try:
                    validation = read_json(validation_path)
                    diff = validation.get("diff") or {}
                    missing = diff.get("missing_from_result") or []
                    extra = diff.get("extra_in_result") or []
                    lines.append(f"- missing orders: `{len(missing)}`")
                    lines.append(f"- extra orders: `{len(extra)}`")
                    if missing[:3]:
                        lines.append(f"- missing sample: `{json.dumps(missing[:3], ensure_ascii=False)}`")
                    if extra[:3]:
                        lines.append(f"- extra sample: `{json.dumps(extra[:3], ensure_ascii=False)}`")
                except Exception:
                    pass
            else:
                artifact = summary.get("paths", {}).get("execution_dir")
                if artifact:
                    lines.append(f"- artifact: [execution_dir]({artifact})")
            lines.append("")
    lines.extend(
        [
            "## Notes",
            "",
            "- `response_format={\"type\":\"json_object\"}`와 Novita top-level `enable_thinking`/`separate_reasoning`을 요청마다 명시했다.",
            "- HTTP 400 계열 파라미터 거부, 인증 오류, image input 미지원, timeout은 `BLOCKED`로 기록하며 파라미터를 제거해 재시도하지 않았다.",
            "- JSON/CSV exact 비교는 status, base_date, issues, 주문 canonical field 기준으로 수행했다.",
        ]
    )
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    write_json(
        artifact_root / "run_summary.json",
        {
            "run_id": run_id,
            "report_path": str(report_path),
            "artifact_root": str(artifact_root),
            "overall": overall,
            "verdict_counts": dict(verdict_counts),
            "profile_summary": profile_summary,
            "summaries": summaries,
        },
    )


def secret_scan(paths: list[Path]) -> list[str]:
    findings: list[str] = []
    patterns = [re.compile(r"sk-[A-Za-z0-9_-]{12,}"), re.compile(r"sk_[A-Za-z0-9_-]{12,}")]
    for path in paths:
        if not path.exists():
            continue
        files = [path] if path.is_file() else [candidate for candidate in path.rglob("*") if candidate.is_file()]
        for file_path in files:
            try:
                text = file_path.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                continue
            for pattern in patterns:
                if pattern.search(text):
                    findings.append(str(file_path))
                    break
    return sorted(set(findings))


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run Novita Qwen3.5 variable-data extraction parameter tests.")
    parser.add_argument("--cycles", type=int, default=2)
    parser.add_argument("--profiles", default="non_thinking,thinking")
    parser.add_argument("--report-root", type=Path, default=REPORT_ROOT)
    parser.add_argument("--base-url", default=None)
    parser.add_argument("--model", default=None)
    parser.add_argument("--timeout", type=int, default=1200)
    parser.add_argument("--case-ids", default=None, help="Comma-separated case ids to run.")
    parser.add_argument("--list-cases", action="store_true")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    cases = [resolve_case(spec) for spec in CASE_SPECS]
    if args.case_ids:
        requested_case_ids = {case_id.strip() for case_id in args.case_ids.split(",") if case_id.strip()}
        known_case_ids = {case["case_id"] for case in cases}
        unknown_case_ids = sorted(requested_case_ids - known_case_ids)
        if unknown_case_ids:
            raise ValueError(f"Unknown case ids: {', '.join(unknown_case_ids)}")
        cases = [case for case in cases if case["case_id"] in requested_case_ids]
    if args.list_cases:
        for case in cases:
            print(f"{case['case_id']}\t{case['company']}\t{case['source']}\t{case['expected_orders']}\t{case['base_date']}")
        return 0

    load_env_file(ROOT / ".env")
    api_key = os.getenv("NOVITA_API_KEY") or os.getenv("LLM_API_KEY")
    if not api_key:
        raise RunBlockedError("NOVITA_API_KEY or LLM_API_KEY is missing")
    base_url = args.base_url or os.getenv("NOVITA_BASE_URL") or os.getenv("LLM_BASE_URL") or DEFAULT_BASE_URL
    endpoint = chat_url(base_url)
    env_model = os.getenv("NOVITA_MODEL") or os.getenv("LLM_MODEL")
    model = args.model or os.getenv("NOVITA_MODEL") or DEFAULT_MODEL
    if not args.model and not os.getenv("NOVITA_MODEL") and env_model and normalize_name(env_model).lower() != DEFAULT_MODEL:
        print(f"warning: ignoring LLM_MODEL={env_model!r}; target model is {DEFAULT_MODEL!r}. Use --model to override.", flush=True)
    selected_profiles = [PROFILES[name.strip()] for name in args.profiles.split(",") if name.strip()]

    now = datetime.now(KST)
    date_dir = args.report_root / now.strftime("%Y%m%d")
    run_id = f"novita_qwen35_variable_data_{now.strftime('%H%M%S')}"
    report_path = date_dir / f"{run_id}.md"
    artifact_root = date_dir / f"{run_id}_artifacts"
    asset_root = artifact_root / "case_assets"
    artifact_root.mkdir(parents=True, exist_ok=True)

    case_manifests: dict[str, dict[str, Any]] = {}
    for case in cases:
        case_manifests[case["case_id"]] = render_assets(case, asset_root)
    write_json(asset_root / "case_manifest.json", {"cases": [{k: str(v) if isinstance(v, Path) else v for k, v in case.items() if k != "spec"} for case in cases]})

    summaries: list[dict[str, Any]] = []
    for cycle in range(1, args.cycles + 1):
        for profile in selected_profiles:
            for case in cases:
                execution_dir = artifact_root / "executions" / f"cycle_{cycle:02d}" / profile.name / case["case_id"]
                summary = run_case_execution(
                    case,
                    case_manifests[case["case_id"]],
                    cycle,
                    profile,
                    execution_dir,
                    endpoint=endpoint,
                    api_key=api_key,
                    model=model,
                    timeout=args.timeout,
                )
                summaries.append(summary)
                print(
                    f"{cycle}/{profile.name}/{case['case_id']}: {summary['verdict']} "
                    f"orders={summary.get('orders')}/{summary.get('expected_orders')} "
                    f"elapsed={summary.get('elapsed_seconds')}",
                    flush=True,
                )
                write_report(
                    report_path,
                    run_id=run_id,
                    artifact_root=artifact_root,
                    cases=cases,
                    profiles=selected_profiles,
                    cycles=args.cycles,
                    endpoint=endpoint,
                    model=model,
                    summaries=summaries,
                )

    findings = secret_scan([report_path, artifact_root])
    write_json(artifact_root / "secret_scan.json", {"findings": findings, "clean": not findings})
    if findings:
        raise RuntimeError(f"API key-like secret pattern found in artifacts: {findings}")

    write_report(
        report_path,
        run_id=run_id,
        artifact_root=artifact_root,
        cases=cases,
        profiles=selected_profiles,
        cycles=args.cycles,
        endpoint=endpoint,
        model=model,
        summaries=summaries,
    )
    print(f"report={report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
